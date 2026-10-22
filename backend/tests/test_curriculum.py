import pytest
from datetime import date
from tests.conftest import TestingSessionLocal
from app.models.all_models import Program, Batch, Student, CurriculumTopic, CurriculumProgress

MENTOR_HEADERS = {"Authorization": "Bearer mock_jwt_token_usr-men-001"}
STUDENT_HEADERS = {"Authorization": "Bearer mock_jwt_token_usr-student-001"}
HR_HEADERS = {"Authorization": "Bearer mock_jwt_token_usr-hr-001"}

@pytest.fixture(autouse=True)
def seed_curriculum_test_data():
    db = TestingSessionLocal()
    prog = Program(
        id="prog-test-curriculum",
        title="Full Stack Engineering",
        code="FSE-TEST",
        duration_weeks=12,
        is_active=True,
        created_by="usr-hr-001"
    )
    db.add(prog)
    db.commit()

    batch = Batch(
        id="batch-curriculum-01",
        program_id="prog-test-curriculum",
        name="Batch Alpha",
        start_date=date(2026, 1, 1),
        end_date=date(2026, 6, 1),
        mentor_id="usr-men-001",
        status="ACTIVE"
    )
    db.add(batch)
    db.commit()

    student = Student(
        id="std-curriculum-01",
        user_id="usr-student-001",
        program_id="prog-test-curriculum",
        batch_id="batch-curriculum-01",
        mentor_id="usr-men-001",
        roll_number="NL-TEST-001",
        enrollment_status="ACTIVE"
    )
    db.add(student)
    db.commit()

    # Seed 4 topics
    t1 = CurriculumTopic(id="top-01", program_id="prog-test-curriculum", module_name="Module 1", title="HTML5 Semantics", order_index=1)
    t2 = CurriculumTopic(id="top-02", program_id="prog-test-curriculum", module_name="Module 1", title="Modern CSS3", order_index=2)
    t3 = CurriculumTopic(id="top-03", program_id="prog-test-curriculum", module_name="Module 2", title="TypeScript Fundamentals", order_index=3)
    t4 = CurriculumTopic(id="top-04", program_id="prog-test-curriculum", module_name="Module 2", title="FastAPI Microservices", order_index=4)
    db.add_all([t1, t2, t3, t4])
    db.commit()
    db.close()


def test_get_program_curriculum(client):
    res = client.get("/api/v1/curriculum/programs/prog-test-curriculum", headers=MENTOR_HEADERS)
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["program_id"] == "prog-test-curriculum"
    assert data["total_topics"] == 4
    assert data["completed_topics"] == 0
    assert data["completion_percentage"] == 0.0
    assert len(data["topics"]) == 4


def test_mentor_update_topic_progress_success(client):
    # Mentor marks topic 1 as completed after an assessment
    payload = {
        "is_completed": True,
        "assessment_name": "Sprint 1 HTML/CSS Assessment",
        "notes": "95% batch scored above 85 points.",
        "batch_id": "batch-curriculum-01"
    }
    update_res = client.put("/api/v1/curriculum/topics/top-01/progress", json=payload, headers=MENTOR_HEADERS)
    assert update_res.status_code == 200, update_res.text
    data = update_res.json()
    assert data["id"] == "top-01"
    assert data["is_completed"] is True
    assert data["assessment_name"] == "Sprint 1 HTML/CSS Assessment"

    # Verify curriculum now shows 1 of 4 completed (25.0%)
    curr_res = client.get("/api/v1/curriculum/programs/prog-test-curriculum?batch_id=batch-curriculum-01", headers=MENTOR_HEADERS)
    assert curr_res.status_code == 200
    curr_data = curr_res.json()
    assert curr_data["completed_topics"] == 1
    assert curr_data["completion_percentage"] == 25.0


def test_student_and_hr_can_view_updated_curriculum(client):
    # Mark topic 1 completed
    payload = {
        "is_completed": True,
        "assessment_name": "Frontend Assessment",
        "notes": "Completed successfully.",
        "batch_id": "batch-curriculum-01"
    }
    client.put("/api/v1/curriculum/topics/top-01/progress", json=payload, headers=MENTOR_HEADERS)

    # Student views curriculum
    stu_res = client.get("/api/v1/curriculum/my-curriculum", headers=STUDENT_HEADERS)
    assert stu_res.status_code == 200
    stu_data = stu_res.json()
    assert stu_data["completed_topics"] == 1
    assert stu_data["topics"][0]["is_completed"] is True
    assert stu_data["topics"][0]["assessment_name"] == "Frontend Assessment"

    # HR views curriculum
    hr_res = client.get("/api/v1/curriculum/programs/prog-test-curriculum?batch_id=batch-curriculum-01", headers=HR_HEADERS)
    assert hr_res.status_code == 200
    hr_data = hr_res.json()
    assert hr_data["completed_topics"] == 1
    assert hr_data["completion_percentage"] == 25.0


def test_student_cannot_update_topic_progress(client):
    # Student attempts to mark topic as completed -> 403 Forbidden
    payload = {
        "is_completed": True,
        "assessment_name": "Self Evaluation",
        "batch_id": "batch-curriculum-01"
    }
    res = client.put("/api/v1/curriculum/topics/top-02/progress", json=payload, headers=STUDENT_HEADERS)
    assert res.status_code == 403


def test_hr_curriculum_import_and_publish_workflow(client):
    # 1. HR imports new curriculum topics for the program
    import_payload = {
        "replace_existing": False,
        "topics": [
            {
                "module_name": "Cloud DevOps",
                "title": "Docker Containerization",
                "description": "Dockerfiles and container lifecycle",
                "order_index": 5
            },
            {
                "module_name": "Cloud DevOps",
                "title": "Kubernetes Orchestration",
                "description": "Pods, Deployments, and Services",
                "order_index": 6
            }
        ]
    }
    import_res = client.post(
        "/api/v1/curriculum/programs/prog-test-curriculum/import",
        json=import_payload,
        headers=HR_HEADERS
    )
    assert import_res.status_code == 200, import_res.text
    import_data = import_res.json()
    assert import_data["total_topics"] == 6  # 4 initial + 2 imported

    # 2. HR unpublishes the curriculum
    unpublish_res = client.put(
        "/api/v1/curriculum/programs/prog-test-curriculum/publish",
        json={"is_published": False},
        headers=HR_HEADERS
    )
    assert unpublish_res.status_code == 200
    assert unpublish_res.json()["is_published"] is False

    # 3. Student views unpublished curriculum -> hidden topics, is_published=False
    student_curr = client.get("/api/v1/curriculum/my-curriculum", headers=STUDENT_HEADERS)
    assert student_curr.status_code == 200
    st_json = student_curr.json()
    assert st_json["is_published"] is False
    assert len(st_json["topics"]) == 0

    # 4. HR still sees all topics even when unpublished
    hr_curr = client.get("/api/v1/curriculum/programs/prog-test-curriculum", headers=HR_HEADERS)
    assert hr_curr.status_code == 200
    hr_json = hr_curr.json()
    assert hr_json["is_published"] is False
    assert len(hr_json["topics"]) == 6

    # 5. HR publishes the curriculum
    publish_res = client.put(
        "/api/v1/curriculum/programs/prog-test-curriculum/publish",
        json={"is_published": True},
        headers=HR_HEADERS
    )
    assert publish_res.status_code == 200
    assert publish_res.json()["is_published"] is True

    # 6. Student can now see all published topics
    student_published = client.get("/api/v1/curriculum/my-curriculum", headers=STUDENT_HEADERS)
    assert student_published.status_code == 200
    st_pub_json = student_published.json()
    assert st_pub_json["is_published"] is True
    assert len(st_pub_json["topics"]) == 6


def test_student_cannot_import_or_publish_curriculum(client):
    # Student attempts to import -> 403
    imp_res = client.post(
        "/api/v1/curriculum/programs/prog-test-curriculum/import",
        json={"replace_existing": False, "topics": [{"module_name": "M", "title": "T", "order_index": 1}]},
        headers=STUDENT_HEADERS
    )
    assert imp_res.status_code == 403

    # Student attempts to publish -> 403
    pub_res = client.put(
        "/api/v1/curriculum/programs/prog-test-curriculum/publish",
        json={"is_published": True},
        headers=STUDENT_HEADERS
    )
    assert pub_res.status_code == 403


def test_parse_curriculum_doc_file_success(client):
    import io
    import docx

    doc = docx.Document()
    doc.add_heading("Module 1: Cloud Architecture", level=1)
    doc.add_paragraph("1. Microservices Basics - Decoupled architectures and container orchestration.")
    doc.add_paragraph("2. API Gateways - Routing, rate limiting, and auth token validation.")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    files = {"file": ("cloud_curriculum.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/api/v1/curriculum/parse-doc", files=files, headers=HR_HEADERS)
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["total_topics"] == 2
    assert data["topics"][0]["module_name"] == "Module 1: Cloud Architecture"
    assert data["topics"][0]["title"] == "Microservices Basics"
    assert "Decoupled architectures" in data["topics"][0]["description"]


def test_parse_curriculum_doc_file_rejects_non_doc(client):
    # Attempting to upload a CSV file instead of doc/docx
    files = {"file": ("curriculum.csv", b"Module,Title,Description\nM1,T1,D1", "text/csv")}
    res = client.post("/api/v1/curriculum/parse-doc", files=files, headers=HR_HEADERS)
    assert res.status_code == 400
    assert "Only Word document files (.docx, .doc) are supported" in res.json()["detail"]


def test_import_curriculum_from_doc_creates_new_program(client):
    import io
    import docx

    doc = docx.Document()
    tbl = doc.add_table(rows=1, cols=3)
    tbl.rows[0].cells[0].text = "Module"
    tbl.rows[0].cells[1].text = "Topic Title"
    tbl.rows[0].cells[2].text = "Description"

    r1 = tbl.add_row().cells
    r1[0].text = "Module 1: DevOps"
    r1[1].text = "Docker Fundamentals"
    r1[2].text = "Containerization and Dockerfile authoring"

    r2 = tbl.add_row().cells
    r2[0].text = "Module 1: DevOps"
    r2[1].text = "Kubernetes Clusters"
    r2[2].text = "Pods, services, and ingress deployment"

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    files = {"file": ("devops_track.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {
        "replace_existing": "true",
        "new_program_title": "DevOps Engineering Specialization",
        "new_program_code": "PRG-DEVOPS-2026"
    }

    res = client.post("/api/v1/curriculum/programs/new/import-doc", files=files, data=data, headers=HR_HEADERS)
    assert res.status_code == 200, res.text
    res_data = res.json()
    assert res_data["program_title"] == "DevOps Engineering Specialization"
    assert res_data["program_code"] == "PRG-DEVOPS-2026"
    assert res_data["total_topics"] == 2
    assert res_data["topics"][0]["title"] == "Docker Fundamentals"
    assert res_data["topics"][1]["title"] == "Kubernetes Clusters"


